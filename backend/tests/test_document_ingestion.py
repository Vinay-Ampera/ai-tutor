import re
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pymupdf
import pandas as pd
import xlwt
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import document_ingestion
from app.services.document_ingestion import ingest_document
from app.services.document_parsers.docx_parser import parse_docx
from app.services.document_parsers.pdf_parser import parse_pdf
from app.services.document_parsers.spreadsheet_parser import parse_spreadsheet


def make_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Cell Biology", fontsize=20, fontname="hebo")
    page.insert_text((72, 105), "Cells are the basic unit of life.", fontsize=11)

    x_positions = (72, 220, 370)
    y_positions = (150, 180, 210)
    for x in x_positions:
        page.draw_line((x, y_positions[0]), (x, y_positions[-1]))
    for y in y_positions:
        page.draw_line((x_positions[0], y), (x_positions[-1], y))
    for row_index, row in enumerate(
        (("Cell", "Role"), ("Nucleus", "Stores DNA")),
    ):
        for column_index, value in enumerate(row):
            page.insert_text(
                (x_positions[column_index] + 5, y_positions[row_index] + 20),
                value,
                fontsize=10,
            )

    result = document.tobytes()
    document.close()
    return result


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Photosynthesis", level=1)
    document.add_paragraph("Plants convert light into chemical energy.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Input"
    table.cell(0, 1).text = "Output"
    table.cell(1, 0).text = "Light"
    table.cell(1, 1).text = "Chemical energy"
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def make_workbook() -> bytes:
    result = BytesIO()
    with pd.ExcelWriter(result, engine="openpyxl") as writer:
        pd.DataFrame(
            {"Planet": ["Earth", "Mars"], "Moons": [1, 2]}
        ).to_excel(writer, sheet_name="Solar System", index=False)
        pd.DataFrame(
            {"State": ["Solid", "Liquid"], "Shape": ["Fixed", "Variable"]}
        ).to_excel(writer, sheet_name="States of Matter", index=False)
    return result.getvalue()


def make_xls_workbook() -> bytes:
    workbook = xlwt.Workbook()
    solar_system = workbook.add_sheet("Solar System")
    states_of_matter = workbook.add_sheet("States of Matter")
    for row_index, row in enumerate(
        (("Planet", "Moons"), ("Earth", 1), ("Mars", 2))
    ):
        for column_index, value in enumerate(row):
            solar_system.write(row_index, column_index, value)
    for row_index, row in enumerate(
        (("State", "Shape"), ("Solid", "Fixed"), ("Liquid", "Variable"))
    ):
        for column_index, value in enumerate(row):
            states_of_matter.write(row_index, column_index, value)
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


class DocumentParserTests(unittest.TestCase):
    def test_pdf_extracts_page_text_heading_and_table_as_markdown(self) -> None:
        markdown = parse_pdf(make_pdf())

        self.assertIn("## Page 1", markdown)
        self.assertIn("### Cell Biology", markdown)
        self.assertIn("Cells are the basic unit of life.", markdown)
        self.assertIn("| Cell | Role |", markdown)
        self.assertIn("| Nucleus | Stores DNA |", markdown)

    def test_docx_extracts_heading_paragraph_and_table_as_markdown(self) -> None:
        markdown = parse_docx(make_docx())

        self.assertIn("# Photosynthesis", markdown)
        self.assertIn("Plants convert light into chemical energy.", markdown)
        self.assertIn("| Input | Output |", markdown)
        self.assertIn("| Light | Chemical energy |", markdown)

    def test_xlsx_preserves_multiple_sheet_names_and_tabular_data(self) -> None:
        markdown = parse_spreadsheet(make_workbook(), "xlsx")

        self.assertIn("## Sheet: Solar System", markdown)
        self.assertIn("| Planet | Moons |", markdown)
        self.assertIn("| Earth | 1 |", markdown)
        self.assertIn("## Sheet: States of Matter", markdown)
        self.assertIn("| Solid | Fixed |", markdown)

    def test_xls_preserves_multiple_sheet_names_and_tabular_data(self) -> None:
        markdown = parse_spreadsheet(make_xls_workbook(), "xls")

        self.assertIn("## Sheet: Solar System", markdown)
        self.assertIn("| Planet | Moons |", markdown)
        self.assertIn("| Earth | 1 |", markdown)
        self.assertIn("## Sheet: States of Matter", markdown)
        self.assertIn("| Solid | Fixed |", markdown)


class DocumentIngestionTests(unittest.TestCase):
    def test_saved_markdown_filename_has_original_stem_and_unique_timestamp(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            with (
                patch.object(document_ingestion, "MARKDOWN_DIRECTORY", directory),
                patch.object(document_ingestion, "BACKEND_ROOT", directory),
            ):
                first = ingest_document("Lecture Notes.docx", make_docx())
                second = ingest_document("Lecture Notes.docx", make_docx())

            filename_pattern = re.compile(
                r"^Lecture_Notes_\d{8}_\d{6}_\d{6}\.md$"
            )
            self.assertRegex(first.markdown_filename, filename_pattern)
            self.assertRegex(second.markdown_filename, filename_pattern)
            self.assertNotEqual(first.markdown_filename, second.markdown_filename)
            self.assertEqual(first.original_filename, "Lecture Notes.docx")
            self.assertEqual(first.file_type, "docx")
            self.assertTrue((directory / first.markdown_filename).is_file())
            self.assertTrue((directory / second.markdown_filename).is_file())
            saved_markdown = (directory / first.markdown_filename).read_text()
            self.assertIn('original_filename: "Lecture Notes.docx"', saved_markdown)
            self.assertIn(f'markdown_filename: "{first.markdown_filename}"', saved_markdown)

    def test_upload_endpoint_returns_metadata_and_rejects_invalid_files(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            with (
                patch.object(document_ingestion, "MARKDOWN_DIRECTORY", directory),
                patch.object(document_ingestion, "BACKEND_ROOT", directory),
                TestClient(app) as client,
            ):
                uploaded = client.post(
                    "/api/documents/ingest",
                    files={
                        "file": (
                            "science.xlsx",
                            make_workbook(),
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet",
                        )
                    },
                )
                unsupported = client.post(
                    "/api/documents/ingest",
                    files={"file": ("notes.txt", b"not supported", "text/plain")},
                )
                corrupt_uploads = {
                    file_type: client.post(
                        "/api/documents/ingest",
                        files={
                            "file": (
                                f"broken.{file_type}",
                                b"not a document",
                                "application/octet-stream",
                            )
                        },
                    )
                    for file_type in ("pdf", "docx", "xls", "xlsx")
                }

            self.assertEqual(uploaded.status_code, 201)
            metadata = uploaded.json()["metadata"]
            self.assertEqual(metadata["original_filename"], "science.xlsx")
            self.assertEqual(metadata["file_type"], "xlsx")
            self.assertTrue(metadata["processed_at"])
            self.assertTrue(metadata["markdown_path"].endswith(
                metadata["markdown_filename"]
            ))
            self.assertTrue((directory / metadata["markdown_filename"]).is_file())
            self.assertEqual(unsupported.status_code, 415)
            self.assertIn("PDF, DOCX, XLS, or XLSX", unsupported.json()["detail"])
            for file_type, response in corrupt_uploads.items():
                with self.subTest(file_type=file_type):
                    self.assertEqual(response.status_code, 422)
                    self.assertIn(file_type.upper(), response.json()["detail"])
