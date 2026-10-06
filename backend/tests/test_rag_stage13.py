import json
import os
import unittest
from io import BytesIO
from datetime import UTC, datetime
from uuid import uuid4
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.graphs.tutor_graph import clear_tutor_session, run_tutor_graph
from app.main import app
from app.graphs.retrieval_nodes import check_document_readiness
from app.services.document_ingestion import (
    DocumentIngestionError,
    DocumentIngestionResult,
)
from app.services.rag.indexing import IndexingResult
from app.services.rag.postgres_store import connect_database
from app.services.document_ingestion import BACKEND_ROOT


class DocumentReadyFlowTests(unittest.TestCase):
    def test_upload_parses_indexes_and_returns_ready_only_after_indexing(self) -> None:
        document_id = uuid4()
        processed_at = datetime.now(UTC).isoformat()
        ingestion = DocumentIngestionResult(
            original_filename="biology.pdf",
            file_type="pdf",
            processed_at=processed_at,
            markdown_filename="biology_20261006_143522_123456.md",
            markdown_path="data/markdown/biology_20261006_143522_123456.md",
        )
        indexed = IndexingResult(
            document_id=document_id,
            original_filename="biology.pdf",
            source_format="pdf",
            markdown_filename=ingestion.markdown_filename,
            markdown_path=ingestion.markdown_path,
            chunk_count=3,
            embedding_dimension=384,
            processing_status="READY",
        )
        with (
            patch(
                "app.main.start_session_document_processing",
                return_value=document_id,
            ) as start_processing,
            patch("app.main.ingest_document", return_value=ingestion) as ingest,
            patch(
                "app.main.index_stage10_markdown",
                return_value=indexed,
            ) as index_markdown,
            TestClient(app) as client,
        ):
            response = client.post(
                "/api/documents/upload",
                data={"session_id": "learner-session"},
                files={"file": ("biology.pdf", b"pdf bytes", "application/pdf")},
            )

        self.assertEqual(response.status_code, 201)
        result = response.json()
        self.assertEqual(result["processing_status"], "READY")
        self.assertEqual(result["chunk_count"], 3)
        self.assertIn("document is ready", result["message"])
        start_processing.assert_called_once_with(
            "learner-session",
            "biology.pdf",
            "pdf",
        )
        ingest.assert_called_once_with("biology.pdf", b"pdf bytes")
        index_markdown.assert_called_once_with(
            ingestion.markdown_filename,
            document_id=document_id,
            ready_status="READY",
        )

    def test_upload_records_failure_instead_of_returning_ready(self) -> None:
        document_id = uuid4()
        with (
            patch(
                "app.main.start_session_document_processing",
                return_value=document_id,
            ),
            patch(
                "app.main.ingest_document",
                side_effect=DocumentIngestionError("The PDF is unreadable."),
            ),
            patch("app.main.mark_session_document_failed") as mark_failed,
            TestClient(app) as client,
        ):
            response = client.post(
                "/api/documents/upload",
                data={"session_id": "learner-session"},
                files={"file": ("broken.pdf", b"bad pdf", "application/pdf")},
            )

        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["detail"], "The PDF is unreadable.")
        mark_failed.assert_called_once_with(document_id, "The PDF is unreadable.")

    def test_session_status_endpoint_reports_empty_and_ready_states(self) -> None:
        with (
            patch("app.main.session_document_status", return_value=None),
            TestClient(app) as client,
        ):
            empty_response = client.get("/api/sessions/empty-session/document")
        self.assertEqual(empty_response.json()["processing_status"], "EMPTY")

        with (
            patch(
                "app.main.session_document_status",
                return_value={
                    "processing_status": "READY",
                    "original_filename": "biology.pdf",
                    "processing_error": None,
                    "markdown_filename": "biology.md",
                },
            ),
            TestClient(app) as client,
        ):
            ready_response = client.get("/api/sessions/ready-session/document")
        self.assertEqual(ready_response.json()["processing_status"], "READY")
        self.assertEqual(ready_response.json()["original_filename"], "biology.pdf")

    def test_graph_returns_processing_message_without_retrieval(self) -> None:
        state = {"session_id": "learner-session"}
        with (
            patch(
                "app.graphs.retrieval_nodes.session_document_status",
                return_value={"processing_status": "PROCESSING"},
            ) as ready_lookup,
            patch("app.graphs.retrieval_nodes.embed_query") as embed_query,
            patch("app.graphs.retrieval_nodes.search_similar_chunks") as search,
        ):
            result = check_document_readiness(state)

        self.assertEqual(result["next_action"], "document_not_ready")
        self.assertIn("still being prepared", result["response"])
        ready_lookup.assert_called_once_with("learner-session")
        embed_query.assert_not_called()
        search.assert_not_called()

    def test_graph_returns_failure_message_without_retrieval(self) -> None:
        with (
            patch(
                "app.graphs.retrieval_nodes.session_document_status",
                return_value={"processing_status": "FAILED"},
            ),
            patch("app.graphs.retrieval_nodes.embed_query") as embed_query,
        ):
            result = check_document_readiness({"session_id": "learner-session"})

        self.assertIn("could not be prepared", result["response"])
        embed_query.assert_not_called()

    def test_real_upload_index_and_session_retrieval(self) -> None:
        if os.getenv("RUN_RAG_POSTGRES_INTEGRATION") != "1":
            self.skipTest("Set RUN_RAG_POSTGRES_INTEGRATION=1 to use PostgreSQL.")

        session_id = f"stage13-{uuid4().hex}"
        document = Document()
        document.add_heading("Cell Biology", level=1)
        document.add_paragraph(
            "Mitochondria produce ATP through cellular respiration. "
            "The nucleus stores the cell's genetic material."
        )
        file_bytes = BytesIO()
        document.save(file_bytes)
        markdown_path = None

        try:
            with TestClient(app) as client:
                uploaded = client.post(
                    "/api/documents/upload",
                    data={"session_id": session_id},
                    files={
                        "file": (
                            "stage13_cells.docx",
                            file_bytes.getvalue(),
                            "application/vnd.openxmlformats-officedocument."
                            "wordprocessingml.document",
                        )
                    },
                )
                self.assertEqual(uploaded.status_code, 201, uploaded.text)
                upload_result = uploaded.json()
                markdown_path = BACKEND_ROOT / upload_result["metadata"]["markdown_path"]
                self.assertEqual(upload_result["processing_status"], "READY")
                self.assertGreater(upload_result["chunk_count"], 0)
                self.assertEqual(upload_result["embedding_dimension"], 384)

                status = client.get(f"/api/sessions/{session_id}/document")
                self.assertEqual(status.json()["processing_status"], "READY")

            with (
                patch(
                    "app.services.gemini.generate_text",
                    side_effect=[
                        json.dumps(
                            {
                                "request_category": "educational",
                                "is_greeting": False,
                                "is_educational": True,
                            }
                        ),
                        "The document says mitochondria produce ATP through "
                        "cellular respiration.",
                    ],
                )
            ):
                answer = run_tutor_graph(
                    "According to the uploaded document, what do mitochondria do?",
                    session_id=session_id,
                )

            self.assertEqual(answer["next_action"], "document_answer")
            self.assertIn("produce ATP", answer["response"])
        finally:
            clear_tutor_session(session_id)
            with connect_database() as connection:
                with connection.cursor() as cursor:
                    cursor.execute(
                        "DELETE FROM rag_documents WHERE session_id = %s",
                        (session_id,),
                    )
            if markdown_path and markdown_path.exists():
                markdown_path.unlink()


if __name__ == "__main__":
    unittest.main()
