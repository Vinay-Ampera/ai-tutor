import json
import os
import re
import tempfile
import unittest
from io import BytesIO
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch
from uuid import uuid4

import numpy as np
from dotenv import load_dotenv
from docx import Document

from app.graphs.tutor_graph import run_tutor_graph
from app.services import document_ingestion
from app.services.rag import embeddings
from app.services.rag.chunking import chunk_markdown
from app.services.rag.embeddings import EMBEDDING_DIMENSION, embed_query, embed_texts
from app.services.rag.indexing import index_stage10_markdown
from app.services.rag.markdown_loader import (
    MarkdownDocumentError,
    MarkdownSource,
    load_stage10_markdown,
)
from app.services.rag import markdown_loader, postgres_store
from app.services.rag.postgres_store import ensure_rag_schema, store_document_chunks
from app.services.rag.postgres_store import (
    connect_database,
    search_similar_chunks,
)


def rough_token_count(text: str) -> int:
    return len(re.findall(r"\w+|[^\w\s]", text))


def stage10_markdown() -> str:
    metadata = {
        "original_filename": "biology_notes.pdf",
        "file_type": "pdf",
        "processed_at": "2026-10-06T10:00:00+00:00",
        "markdown_filename": "biology_notes_20261006_100000_123456.md",
        "markdown_path": "data/markdown/biology_notes_20261006_100000_123456.md",
    }
    return (
        "---\n"
        + "\n".join(
            f"{key}: {json.dumps(value)}" for key, value in metadata.items()
        )
        + "\n---\n\n"
        + "## Page 2\n\n"
        + "### Photosynthesis\n\n"
        + "Plants use sunlight to convert water and carbon dioxide into glucose. "
        + "The process releases oxygen as a by-product.\n\n"
        + "| Input | Output |\n"
        + "| --- | --- |\n"
        + "| Water | Oxygen |\n"
        + "| Carbon dioxide | Glucose |\n"
    )


class ChunkingTests(unittest.TestCase):
    def test_chunks_preserve_page_heading_and_table_structure(self) -> None:
        chunks = chunk_markdown(
            stage10_markdown().split("\n---\n\n", maxsplit=1)[1],
            max_tokens=36,
            token_counter=rough_token_count,
        )

        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(all(chunk.page_number == 2 for chunk in chunks))
        self.assertTrue(all(
            "Page 2 > Photosynthesis" in (chunk.heading or "")
            for chunk in chunks
        ))
        self.assertIn("Plants use sunlight", chunks[0].content)
        table_chunks = [chunk.content for chunk in chunks if "| Input | Output |" in chunk.content]
        self.assertTrue(table_chunks)
        self.assertIn("| --- | --- |", table_chunks[0])
        self.assertIn("| Water | Oxygen |", table_chunks[0])

    def test_chunks_split_oversized_sections_by_sentences_and_keep_context(self) -> None:
        markdown = (
            "## Sheet: Chemistry\n\n"
            "### Acids\n\n"
            "Acids donate protons in a reaction. Bases accept protons in a reaction. "
            "The pH scale describes acidity."
        )
        chunks = chunk_markdown(
            markdown,
            max_tokens=14,
            token_counter=rough_token_count,
        )

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.sheet_name == "Chemistry" for chunk in chunks))
        self.assertTrue(all(
            "Sheet: Chemistry > Acids" in (chunk.heading or "")
            for chunk in chunks
        ))
        self.assertTrue(all(rough_token_count(chunk.content) <= 14 for chunk in chunks))

    def test_invalid_chunk_size_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "positive integer"):
            chunk_markdown("Text", max_tokens=0)


class MarkdownLoaderTests(unittest.TestCase):
    def test_loads_stage10_frontmatter_and_markdown_content(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            filename = "biology_notes_20261006_100000_123456.md"
            (directory / filename).write_text(stage10_markdown(), encoding="utf-8")
            with patch.object(markdown_loader, "MARKDOWN_DIRECTORY", directory):
                source = load_stage10_markdown(filename)

        self.assertEqual(source.original_filename, "biology_notes.pdf")
        self.assertEqual(source.source_format, "pdf")
        self.assertEqual(source.uploaded_at, datetime(2026, 10, 6, 10, 0, tzinfo=UTC))
        self.assertIn("## Page 2", source.content)

    def test_rejects_paths_outside_generated_markdown_directory(self) -> None:
        with self.assertRaisesRegex(MarkdownDocumentError, "filename"):
            load_stage10_markdown("../private.md")

    def test_rejects_missing_stage10_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            (directory / "plain.md").write_text("# No metadata", encoding="utf-8")
            with patch.object(markdown_loader, "MARKDOWN_DIRECTORY", directory):
                with self.assertRaisesRegex(MarkdownDocumentError, "frontmatter"):
                    load_stage10_markdown("plain.md")


class EmbeddingTests(unittest.TestCase):
    def test_document_and_query_embedding_use_same_model_and_dimension(self) -> None:
        model = MagicMock()
        model.encode.side_effect = lambda texts, **kwargs: np.tile(
            np.eye(1, EMBEDDING_DIMENSION, dtype=np.float32),
            (len(texts), 1),
        )
        with patch.object(embeddings, "_load_model", return_value=model):
            document_vectors = embed_texts(["A document chunk"])
            query_vector = embed_query("A sample search question")

        self.assertEqual(len(document_vectors), 1)
        self.assertEqual(len(document_vectors[0]), EMBEDDING_DIMENSION)
        self.assertEqual(len(query_vector), EMBEDDING_DIMENSION)
        self.assertEqual(model.encode.call_count, 2)
        self.assertTrue(all(
            call.kwargs["normalize_embeddings"]
            for call in model.encode.call_args_list
        ))
        self.assertEqual(embeddings.MODEL_NAME, "BAAI/bge-small-en-v1.5")

    def test_empty_embedding_inputs_are_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "empty text"):
            embed_texts([""])
        with self.assertRaisesRegex(ValueError, "must not be empty"):
            embed_query(" ")


class PostgresStoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self.cursor = MagicMock()
        self.cursor.__enter__.return_value = self.cursor
        self.connection = MagicMock()
        self.connection.__enter__.return_value = self.connection
        self.connection.cursor.return_value = self.cursor

    def test_database_url_alone_is_not_used_for_connection(self) -> None:
        with (
            patch.dict(os.environ, {"DATABASE_URL": "unused"}, clear=True),
            patch("dotenv.load_dotenv"),
            self.assertRaisesRegex(
                postgres_store.DatabaseConfigurationError,
                "DB_HOST.*DB_PASSWORD",
            ),
        ):
            postgres_store.connect_database()

    def test_postgres_connection_uses_separate_settings_and_literal_password(self) -> None:
        settings = {
            "DB_HOST": "localhost",
            "DB_PORT": "5432",
            "DB_NAME": "ai_tutor",
            "DB_USER": "test_user",
            "DB_PASSWORD": "a@password",
        }
        with (
            patch.dict(os.environ, settings, clear=True),
            patch("app.services.rag.postgres_store.psycopg.connect") as connect,
        ):
            connect.return_value = self.connection
            connection = postgres_store.connect_database()

        self.assertIs(connection, self.connection)
        connect.assert_called_once_with(
            host="localhost",
            port=5432,
            dbname="ai_tutor",
            user="test_user",
            password="a@password",
            row_factory=postgres_store.dict_row,
        )

    def test_schema_uses_bge_dimension_foreign_key_and_hnsw(self) -> None:
        self.cursor.fetchone.return_value = {"exists": True}
        with patch("app.services.rag.postgres_store.register_vector"):
            ensure_rag_schema(self.connection)

        executed_sql = "\n".join(
            call.args[0] for call in self.cursor.execute.call_args_list
        )
        self.assertIn(f"vector({EMBEDDING_DIMENSION})", executed_sql)
        self.assertIn("REFERENCES rag_documents(id) ON DELETE CASCADE", executed_sql)
        self.assertIn("USING hnsw", executed_sql)
        self.assertIn("vector_cosine_ops", executed_sql)

    def test_latest_document_lookup_only_selects_indexed_records(self) -> None:
        self.cursor.fetchone.return_value = None
        with patch(
            "app.services.rag.postgres_store.connect_database",
            return_value=self.connection,
        ):
            result = postgres_store.latest_indexed_document()

        self.assertIsNone(result)
        query = self.cursor.execute.call_args.args[0]
        self.assertIn("processing_status IN ('INDEXED', 'READY')", query)
        self.assertIn("ORDER BY indexed_at DESC", query)

    def test_stores_and_verifies_document_chunks_and_metadata(self) -> None:
        self.cursor.fetchone.side_effect = [
            {"exists": True},
            None,
            {
                "chunk_count": 1,
                "minimum_dimension": EMBEDDING_DIMENSION,
                "maximum_dimension": EMBEDDING_DIMENSION,
            },
        ]
        source = MarkdownSource(
            original_filename="biology.pdf",
            source_format="pdf",
            uploaded_at=datetime(2026, 10, 6, tzinfo=UTC),
            markdown_filename="biology_20261006.md",
            markdown_path="data/markdown/biology_20261006.md",
            content="# Biology",
        )
        vectors = [[0.0] * (EMBEDDING_DIMENSION - 1) + [1.0]]

        with (
            patch("app.services.rag.postgres_store.connect_database", return_value=self.connection),
            patch("app.services.rag.postgres_store.register_vector"),
        ):
            document_id, stored_count = store_document_chunks(
                source,
                [("# Biology\n\nCell notes", {"heading": "Biology"})],
                vectors,
            )

        self.assertIsInstance(document_id, type(uuid4()))
        self.assertEqual(stored_count, 1)
        self.cursor.executemany.assert_called_once()
        inserted_values = self.cursor.executemany.call_args.args[1]
        self.assertEqual(inserted_values[0][1], document_id)
        self.assertEqual(inserted_values[0][2], 0)
        self.assertEqual(inserted_values[0][3], "# Biology\n\nCell notes")
        self.assertEqual(inserted_values[0][4].obj, {"heading": "Biology"})
        self.assertEqual(len(inserted_values[0][5].to_list()), EMBEDDING_DIMENSION)

    def test_rejects_embeddings_with_the_wrong_dimension(self) -> None:
        source = MarkdownSource(
            original_filename="biology.pdf",
            source_format="pdf",
            uploaded_at=datetime(2026, 10, 6, tzinfo=UTC),
            markdown_filename="biology.md",
            markdown_path="data/markdown/biology.md",
            content="# Biology",
        )
        with self.assertRaisesRegex(ValueError, "384 dimensions"):
            store_document_chunks(source, [("chunk", {})], [[0.1]])


class IndexingWorkflowTests(unittest.TestCase):
    def test_stage10_markdown_chunks_embeds_and_persists_in_order(self) -> None:
        source = MarkdownSource(
            original_filename="biology.pdf",
            source_format="pdf",
            uploaded_at=datetime(2026, 10, 6, tzinfo=UTC),
            markdown_filename="biology.md",
            markdown_path="data/markdown/biology.md",
            content="## Page 2\n\n### Cells\n\nCells are living units.",
        )
        document_id = uuid4()
        events: list[str] = []

        def embed(texts: list[str]) -> list[list[float]]:
            events.append("embed")
            return [[0.0] * (EMBEDDING_DIMENSION - 1) + [1.0] for _ in texts]

        def store(*args: object) -> tuple[object, int]:
            events.append("store")
            return document_id, len(args[1])

        with (
            patch(
                "app.services.rag.indexing.load_stage10_markdown",
                return_value=source,
            ),
            patch(
                "app.services.rag.indexing.count_embedding_tokens",
                side_effect=rough_token_count,
            ),
            patch("app.services.rag.indexing.embed_texts", side_effect=embed),
            patch(
                "app.services.rag.indexing.store_document_chunks",
                side_effect=store,
            ) as store_mock,
        ):
            result = index_stage10_markdown("biology.md")

        self.assertEqual(events, ["embed", "store"])
        self.assertEqual(result.document_id, document_id)
        self.assertEqual(result.chunk_count, 1)
        self.assertEqual(result.embedding_dimension, EMBEDDING_DIMENSION)
        self.assertEqual(result.processing_status, "INDEXED")
        stored_chunks = store_mock.call_args.args[1]
        self.assertIn("Page 2 > Cells", stored_chunks[0][0])
        self.assertEqual(stored_chunks[0][1]["page_number"], 2)


class PostgresEndToEndTests(unittest.TestCase):
    def test_real_markdown_index_embedding_storage_and_vector_search(self) -> None:
        if os.getenv("RUN_RAG_POSTGRES_INTEGRATION") != "1":
            self.skipTest("Set RUN_RAG_POSTGRES_INTEGRATION=1 to use PostgreSQL.")
        backend_directory = Path(__file__).resolve().parents[1]
        load_dotenv(backend_directory / ".env")
        database_settings = (
            "DB_HOST",
            "DB_PORT",
            "DB_NAME",
            "DB_USER",
            "DB_PASSWORD",
        )
        if any(not os.getenv(setting, "") for setting in database_settings):
            self.skipTest(
                "Set DB_HOST, DB_PORT, DB_NAME, DB_USER, and DB_PASSWORD "
                "in backend/.env for PostgreSQL integration."
            )

        with tempfile.TemporaryDirectory() as temporary_directory:
            backend_directory = Path(temporary_directory)
            markdown_directory = backend_directory / "data" / "markdown"
            docx = Document()
            docx.add_heading("Verification section", level=1)
            docx.add_paragraph(
                "BGE stores this verification phrase in a searchable document."
            )
            docx_content = BytesIO()
            docx.save(docx_content)
            result = None
            try:
                with (
                    patch.object(
                        document_ingestion,
                        "MARKDOWN_DIRECTORY",
                        markdown_directory,
                    ),
                    patch.object(
                        document_ingestion,
                        "BACKEND_ROOT",
                        backend_directory,
                    ),
                    patch.object(
                        markdown_loader,
                        "MARKDOWN_DIRECTORY",
                        markdown_directory,
                    ),
                ):
                    stage10_result = document_ingestion.ingest_document(
                        "stage11_verification.docx",
                        docx_content.getvalue(),
                    )
                    result = index_stage10_markdown(
                        stage10_result.markdown_filename
                    )
                matches = search_similar_chunks(
                    embed_query("Which phrase is stored in the verification document?"),
                    document_id=result.document_id,
                    limit=1,
                )
                with patch(
                    "app.services.gemini.generate_text",
                    side_effect=[
                        json.dumps(
                            {
                                "request_category": "educational",
                                "is_greeting": False,
                                "is_educational": True,
                            }
                        ),
                        "The document says BGE stores the verification phrase.",
                    ],
                ):
                    graph_result = run_tutor_graph(
                        "According to the uploaded document, what does BGE store?",
                        session_id=f"stage12-{uuid4().hex}",
                    )
                self.assertEqual(result.embedding_dimension, EMBEDDING_DIMENSION)
                self.assertEqual(result.processing_status, "INDEXED")
                self.assertGreaterEqual(result.chunk_count, 1)
                self.assertEqual(len(matches), 1)
                self.assertIn("BGE stores this verification phrase", matches[0]["content"])
                self.assertEqual(matches[0]["metadata"]["heading"], "Verification section")
                self.assertIn("BGE stores the verification phrase", graph_result["response"])
                self.assertEqual(graph_result["document_id"], result.document_id)
                self.assertTrue(graph_result["retrieved_chunks"])
            finally:
                if result is not None:
                    with connect_database() as connection:
                        with connection.cursor() as cursor:
                            cursor.execute(
                                "DELETE FROM rag_documents WHERE id = %s",
                                (result.document_id,),
                            )
